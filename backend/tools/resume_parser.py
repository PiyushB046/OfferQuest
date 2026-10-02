"""PDF/DOCX resume → text + best-effort structured fields.

Parsing messy resumes is never perfect, so the result only pre-fills the intake form; the student
corrects it there (spec §19).
"""
from __future__ import annotations

import io
import re

from .lexicon import find_skills

HEADINGS = {
    "summary": r"(career |professional )?(summary|objective|profile)|about( me)?",
    "education": r"education(al)?( background| qualifications?)?|academics?|academic (background|details)|qualifications?",
    "skills": r"(technical |key |core )?skills( (and|&) (tools|technologies|interests))?|technologies|tech stack|technical (proficiency|expertise)",
    "projects": r"(academic |personal |key |major |notable )?projects?( work)?",
    "experience": r"(work |professional |internship |industry )?experience|internships?|employment( history)?|work history",
    "achievements": r"achievements?|awards?( (and|&) (honou?rs|achievements))?|certifications?|accomplishments|"
                    r"extra.?curriculars?( activities)?|positions? of responsibility|honou?rs|courses?|coursework",
}
BRANCHES = ["Computer Science and Engineering", "Computer Science", "Computer Engineering", "Information Technology",
            "Artificial Intelligence and Data Science", "Artificial Intelligence", "Data Science",
            "Electronics and Telecommunication", "Electronics and Communication", "Electronics", "Electrical",
            "Mechanical", "Civil", "Statistics", "Mathematics", "Physics", "Commerce"]


def _page_text(page) -> str:
    """Text in reading order. Two-column resumes are read column by column, not line by line across."""
    from collections import Counter
    blocks = [b for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
    width = page.rect.width
    # A second column shows up as several blocks sharing a left edge well inside the page.
    starts = Counter(round(b[0] / 10) * 10 for b in blocks if b[0] > width * 0.25)
    split, count = starts.most_common(1)[0] if starts else (0, 0)
    left = [b for b in blocks if b[2] <= split + 6]
    if count < 3 or len(left) < 2:
        return "\n".join(b[4] for b in sorted(blocks, key=lambda b: (round(b[1]), b[0])))
    right = [b for b in blocks if b[0] >= split - 6]
    banner = [b for b in blocks if b not in left and b not in right]      # e.g. the name, spanning both columns
    order = sorted(banner, key=lambda b: b[1]) + sorted(left, key=lambda b: b[1]) + sorted(right, key=lambda b: b[1])
    return "\n".join(b[4] for b in order)


def _unwrap(text: str) -> str:
    """Re-joins lines that the PDF wrapped mid-sentence, so one bullet is one line again."""
    out: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        prev = out[-1] if out else ""
        bare = re.sub(r"[^a-z &-]", "", line.lower()).strip()
        heading = any(re.fullmatch(p, bare) for p in HEADINGS.values()) and len(line) < 45
        continues = line[0].islower() or line[0].isdigit() or prev.endswith((",", "-", "–", "/", " and", " with", " in"))
        contact = re.match(r"(https?://|www\.|[\w.+-]+@|\+?\d[\d\s-]{7,}|[\w-]+\.(com|in|io|dev)/)", line)
        if prev and continues and not heading and not _is_bullet(line) and (not contact or prev.endswith("/")):
            out[-1] = prev + ("" if prev.endswith("/") else " ") + line
        else:
            out.append(line)
    return "\n".join(out)


def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        import pymupdf as fitz
        with fitz.open(stream=data, filetype="pdf") as doc:
            return _unwrap("\n".join(_page_text(page) for page in doc))
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        cells = [c.text for t in d.tables for r in t.rows for c in r.cells]   # many resumes are laid out in tables
        return "\n".join([p.text for p in d.paragraphs] + cells)
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Please upload a PDF, DOCX or TXT file.")


def _sections(text: str) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        bare = re.sub(r"[^a-z &-]", "", line.lower()).strip()
        key = next((k for k, pat in HEADINGS.items() if re.fullmatch(pat, bare)), None)
        if key and len(line) < 45:
            current = key
            out.setdefault(current, [])
        else:
            out.setdefault(current, []).append(line)
    return out


def _is_bullet(line: str) -> bool:
    return bool(re.match(r"^[•\-–*·▪◦●]\s*", line))


def _strip_bullet(line: str) -> str:
    return re.sub(r"^[•\-–*·▪◦●]\s*", "", line).strip()


def _projects(lines: list[str]) -> list[dict]:
    projects: list[dict] = []
    sep = r"\s[|–-]\s|\s\("
    for line in lines:
        head = re.split(sep, line)[0].strip()
        # A wrapped bullet also arrives as a plain line: a lone word with no separator is not a title.
        titled = re.search(sep, line) or len(head.split()) >= 2 or not projects
        if not _is_bullet(line) and len(head) < 70 and titled and (not projects or projects[-1]["_body"]):
            link = re.search(r"(https?://\S+|github\.com/\S+)", line)
            projects.append({"name": head, "_body": [], "link": link.group(1) if link else None,
                             "_title_line": line})
        elif projects:
            projects[-1]["_body"].append(_strip_bullet(line))
    out = []
    for p in projects:
        body = ". ".join(b.rstrip(".") for b in p["_body"])
        link = p["link"] or (m.group(1) if (m := re.search(r"(https?://\S+|github\.com/\S+)", body)) else None)
        out.append({"name": p["name"], "description": body,
                    "tech_stack": find_skills(re.sub(r"(https?://)?\S+\.(com|io|dev|app|in)/\S*", " ", p["_title_line"] + " " + body)),
                    "link": link,
                    "metrics": []})
    return out


def _name(header: list[str]) -> str:
    for line in header[:4]:
        words = line.split()
        if (1 < len(words) <= 4 and not re.search(r"[@\d|:/]|resume|curriculum|vitae", line, re.I)
                and all(w[0].isupper() for w in words if w[0].isalpha())):
            return line.title() if line.isupper() else line
    return ""


def parse_resume(text: str, today_year: int | None = None) -> dict:
    from datetime import date
    year_now = today_year or date.today().year
    sec = _sections(text)
    header = sec.get("header", [])
    edu_lines = sec.get("education", [])
    edu = " ".join(edu_lines)
    cgpa = re.search(r"(?:cgpa|gpa|cpi)\D{0,10}(\d(?:\.\d{1,2})?)", text, re.I)
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text)
    github = re.search(r"github\.com/[\w-]+", text, re.I)
    linkedin = re.search(r"linkedin\.com/in/[\w-]+", text, re.I)
    leetcode = re.search(r"leetcode\.com/(?:u/)?[\w-]+", text, re.I)
    degree = re.search(r"\b(B\.?\s?Tech|B\.?E\.?|B\.?Sc|BCA|BBA|B\.?Com|M\.?Tech|MCA|M\.?Sc|MBA)\b", edu or text, re.I)
    years = [int(y) for y in re.findall(r"\b(20[1-4]\d)\b", edu)]
    grad = max(years) if years else None
    college = next((re.split(r"\s[|–-]\s|\(", line)[0].strip(" ,") for line in edu_lines
                    if re.search(r"college|institute|university|school of|academy|\bIIT\b|\bNIT\b|\bIIIT\b", line, re.I)), "")
    branch = next((b for b in BRANCHES if re.search(re.escape(b), edu, re.I)), "")
    experience = []
    for line in sec.get("experience", []):
        if not _is_bullet(line) and len(line) < 110 and (not experience or experience[-1]["highlights"] or not experience[-1]["role"]):
            dates = re.search(r"((?:[A-Z][a-z]{2,8}\.? )?20\d{2})\s*[-–to]+\s*((?:[A-Z][a-z]{2,8}\.? )?20\d{2}|present|current)", line, re.I)
            core = line.replace(dates.group(0), "") if dates else line
            parts = [p.strip(" ()") for p in re.split(r"\s[|–-]\s|,|\sat\s", core) if p.strip(" ()")]
            if not parts:
                continue
            if len(parts) > 1 and re.search(r"intern|engineer|developer|analyst|trainee|assistant|lead|member", parts[0], re.I):
                parts[0], parts[1] = parts[1], parts[0]          # "Role, Company" order
            experience.append({"organisation": parts[0], "role": parts[1] if len(parts) > 1 else "",
                               "start": dates.group(1) if dates else "",
                               "end": None if not dates or dates.group(2).lower() in ("present", "current") else dates.group(2),
                               "highlights": []})
        elif experience:
            experience[-1]["highlights"].append(_strip_bullet(line))
    links = {}
    if email:
        links["email"] = email.group(0)
    if github:
        links["github"] = "https://" + github.group(0)
    if linkedin:
        links["linkedin"] = "https://" + linkedin.group(0)
    if leetcode:
        links["leetcode"] = "https://" + leetcode.group(0)
    fields = {
        "name": _name(header),
        "college": college,
        "degree": degree.group(1) if degree else "",
        "branch": branch,
        "graduation_date": str(grad) if grad else "",
        # A 4-year degree ending in `grad`: someone graduating next year is in final year.
        "year": max(1, min(4, 4 - (grad - year_now) + (0 if date.today().month >= 7 else -1) + 1)) if grad and grad >= year_now else 4 if grad else 3,
        "cgpa": float(cgpa.group(1)) if cgpa else None,
        "skills": find_skills(" ".join(sec.get("skills", [])) or text),
        "projects": _projects(sec.get("projects", [])),
        "experience": experience,
        "achievements": [_strip_bullet(line) for line in sec.get("achievements", []) if len(line) > 3][:8],
        "links": links,
    }
    return fields


def parse_report(fields: dict, text: str) -> list[str]:
    """Plain-language notes for the student about what was and wasn't found."""
    notes = []
    if not fields["name"]:
        notes.append("I couldn't find your name at the top. Please type it.")
    if not fields["projects"]:
        notes.append("I found no Projects section. Add your projects below: they matter most for tailoring.")
    if not fields["skills"]:
        notes.append("I found no skills I recognise. Add them as a comma-separated list.")
    if any(not p["tech_stack"] for p in fields["projects"]):
        notes.append("Some projects have no tech listed. Add what you used.")
    if not any(re.search(r"\d", p["description"]) for p in fields["projects"]) and fields["projects"]:
        notes.append("None of your projects mention a number. If you have real results (users, accuracy, speed), add them.")
    if len(text) > 6500:
        notes.append("This looks longer than one page. Recruiters prefer one page for campus roles.")
    return notes


def merge_llm_fields(fields: dict, llm: dict, text: str) -> dict:
    """Take a model's reading of the resume only where the resume text backs it up."""
    low = text.lower()
    seen = lambda v: isinstance(v, str) and v.strip() and v.strip().lower() in low  # noqa: E731
    out = dict(fields)
    for key in ("name", "college", "degree", "branch"):
        if seen(llm.get(key)):
            out[key] = llm[key].strip().title() if key == "name" and llm[key].isupper() else llm[key].strip()
    if isinstance(llm.get("graduation_date"), (str, int)) and str(llm["graduation_date"]) in text:
        out["graduation_date"] = str(llm["graduation_date"])
    projects = []
    for p in llm.get("projects") or []:
        if not isinstance(p, dict) or not seen(p.get("name")):
            continue
        desc = str(p.get("description") or "")
        tech = [t for t in (p.get("tech_stack") or []) if isinstance(t, str) and t.lower() in low]
        metrics = [m for m in (p.get("metrics") or []) if isinstance(m, str) and all(n in text for n in re.findall(r"\d+(?:\.\d+)?", m))]
        link = p.get("link") if isinstance(p.get("link"), str) and p["link"].lower() in low else None
        projects.append({"name": p["name"].strip(), "description": desc[:900], "tech_stack": tech or find_skills(desc),
                         "link": link, "metrics": metrics[:4]})
    if projects:
        out["projects"] = projects
    extra = [s for s in (llm.get("skills") or []) if isinstance(s, str) and len(s) < 40 and s.lower() in low]
    known = {s.lower() for s in out["skills"]}
    out["skills"] = out["skills"] + [s for s in extra if s.lower() not in known]
    return out
