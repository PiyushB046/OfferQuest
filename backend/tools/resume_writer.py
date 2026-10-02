"""Structured resume → plain text, DOCX and PDF."""
from __future__ import annotations

from html import escape
from pathlib import Path

from ..models import StudentProfile


def split_sentences(text: str) -> list[str]:
    import re
    parts = re.split(r"(?<=[.!?])\s+|\n+|;\s+", text.strip())
    return [p.strip().rstrip(".") for p in parts if len(p.strip()) > 3]


def baseline_resume(st: StudentProfile) -> dict:
    """The student's facts laid out as a resume, before any tailoring."""
    edu = ", ".join(x for x in [f"{st.degree} {st.branch}".strip(), st.college] if x)
    if st.graduation_date:
        edu += f" (graduating {st.graduation_date})"
    if st.cgpa:
        edu += f" · CGPA {st.cgpa}"
    return {
        "name": st.name, "links": dict(st.links), "summary": "", "education": edu,
        "skills": list(st.skills),
        "projects": [{"name": p.name, "tech_stack": list(p.tech_stack), "link": p.link,
                      "bullets": [{"text": s, "orig": s} for s in split_sentences(p.description)]}
                     for p in st.projects],
        "experience": [{"organisation": e.organisation, "role": e.role, "start": e.start, "end": e.end,
                        "bullets": [{"text": h, "orig": h} for h in e.highlights]}
                       for e in st.experience],
        "achievements": list(st.achievements),
    }


def render_text(r: dict) -> tuple[str, str]:
    """Returns (top, full): `top` is the summary + skills block a recruiter reads first."""
    top = "\n".join(x for x in [r["summary"], "Skills: " + ", ".join(r["skills"][:10])] if x)
    lines = [r["name"], " | ".join(r["links"].values()), r["summary"], "EDUCATION", r["education"],
             "SKILLS", ", ".join(r["skills"]), "PROJECTS"]
    for p in r["projects"]:
        lines.append(f"{p['name']} | {', '.join(p['tech_stack'])} {p.get('link') or ''}")
        lines += [f"- {b['text']}" for b in p["bullets"]]
    if r["experience"]:
        lines.append("EXPERIENCE")
        for e in r["experience"]:
            lines.append(f"{e['role']}, {e['organisation']} ({e['start']} - {e['end'] or 'present'})")
            lines += [f"- {b['text']}" for b in e["bullets"]]
    if r["achievements"]:
        lines.append("ACHIEVEMENTS")
        lines += [f"- {a}" for a in r["achievements"]]
    return top, "\n".join(x for x in lines if x)


def write_docx(r: dict, path: Path) -> None:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_heading(r["name"], level=0)
    if r["links"]:
        d.add_paragraph(" | ".join(r["links"].values()))
    if r["summary"]:
        d.add_paragraph(r["summary"])
    d.add_heading("Education", level=2)
    d.add_paragraph(r["education"])
    d.add_heading("Skills", level=2)
    d.add_paragraph(", ".join(r["skills"]))
    d.add_heading("Projects", level=2)
    for p in r["projects"]:
        para = d.add_paragraph()
        para.add_run(p["name"]).bold = True
        para.add_run(f"  |  {', '.join(p['tech_stack'])}" + (f"  |  {p['link']}" if p.get("link") else ""))
        for b in p["bullets"]:
            d.add_paragraph(b["text"], style="List Bullet")
    if r["experience"]:
        d.add_heading("Experience", level=2)
        for e in r["experience"]:
            para = d.add_paragraph()
            para.add_run(f"{e['role']}, {e['organisation']}").bold = True
            para.add_run(f"  ({e['start']} - {e['end'] or 'present'})")
            for b in e["bullets"]:
                d.add_paragraph(b["text"], style="List Bullet")
    if r["achievements"]:
        d.add_heading("Achievements", level=2)
        for a in r["achievements"]:
            d.add_paragraph(a, style="List Bullet")
    d.save(str(path))


def _html(r: dict) -> str:
    h = [f"<h1>{escape(r['name'])}</h1>"]
    if r["links"]:
        h.append(f"<p class='links'>{escape(' | '.join(r['links'].values()))}</p>")
    if r["summary"]:
        h.append(f"<p>{escape(r['summary'])}</p>")
    h.append(f"<h2>Education</h2><p>{escape(r['education'])}</p>")
    h.append(f"<h2>Skills</h2><p>{escape(', '.join(r['skills']))}</p><h2>Projects</h2>")
    for p in r["projects"]:
        meta = ", ".join(p["tech_stack"]) + (f" | {p['link']}" if p.get("link") else "")
        h.append(f"<p><b>{escape(p['name'])}</b> | {escape(meta)}</p><ul>"
                 + "".join(f"<li>{escape(b['text'])}</li>" for b in p["bullets"]) + "</ul>")
    if r["experience"]:
        h.append("<h2>Experience</h2>")
        for e in r["experience"]:
            h.append(f"<p><b>{escape(e['role'])}, {escape(e['organisation'])}</b> "
                     f"({escape(e['start'])} - {escape(e['end'] or 'present')})</p><ul>"
                     + "".join(f"<li>{escape(b['text'])}</li>" for b in e["bullets"]) + "</ul>")
    if r["achievements"]:
        h.append("<h2>Achievements</h2><ul>"
                 + "".join(f"<li>{escape(a)}</li>" for a in r["achievements"]) + "</ul>")
    return "".join(h)


CSS = """body{font-family:sans-serif;font-size:10pt;color:#222}
h1{font-size:20pt;margin:0 0 2pt 0}h2{font-size:11pt;margin:10pt 0 2pt 0;color:#1d4e89;
border-bottom:1px solid #999}p{margin:2pt 0}ul{margin:2pt 0}.links{color:#555;font-size:9pt}"""


def write_pdf(r: dict, path: Path) -> None:
    import pymupdf as fitz
    doc = fitz.open()
    page = doc.new_page()  # A4
    page.insert_htmlbox(fitz.Rect(42, 40, page.rect.width - 42, page.rect.height - 40), _html(r), css=CSS)
    doc.save(str(path))
    doc.close()
